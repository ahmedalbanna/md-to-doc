import os
import tempfile
from markdown_it.token import Token
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .models import DocumentMetadata, ConvertOptions
from .styles import CoverStyles, NumberFormats
from .parsers.frontmatter import parse_front_matter
from .parsers.markdown_parser import MarkdownParser
from .builders.document_builder import DocumentBuilder
from .builders.section_builder import SectionBuilder
from .builders.page_break import is_page_break_token, add_page_break, clean_page_break_text
from .builders.paragraph_builder import ParagraphBuilder, detect_text_direction
from .builders.list_builder import ListBuilder
from .builders.table_builder import TableBuilder
from .builders.image_builder import ImageBuilder
from .builders.code_builder import CodeBuilder
from .builders.diagram_converter import DiagramConverter
from .builders.cover_builder import CoverBuilder
from .builders.numbering_builder import NumberingBuilder
from .builders.header_footer_builder import HeaderFooterBuilder
from .builders.toc_builder import TOCBuilder
from .builders.paragraph_builder import convert_to_arabic_indic


class MarkdownToDocxConverter:
    def __init__(self, metadata: DocumentMetadata = None, options: ConvertOptions = None):
        self.metadata = metadata or DocumentMetadata()
        self.options = options or ConvertOptions()
        self._apply_options()
        self.doc = None
        self.parser = MarkdownParser()
        self.figure_counter = {"figure": 0, "table": 0}
        self.diagram_converter = None

    def _apply_options(self):
        opts = self.options
        if opts.dir:
            self.metadata.dir = opts.dir
        if opts.lang:
            self.metadata.lang = opts.lang
        if opts.font_arabic:
            self.metadata.fonts.arabic = opts.font_arabic
        if opts.font_latin:
            self.metadata.fonts.latin = opts.font_latin
        if opts.cover_style:
            self.metadata.cover_style = opts.cover_style
        if opts.page_number_format:
            self.metadata.page_numbers.body = opts.page_number_format
        self.metadata.heading_numbering.enabled = opts.heading_numbering
        self.metadata.toc.enabled = opts.toc
        self.metadata.lof.enabled = opts.lof
        self.metadata.lot.enabled = opts.lot
        self.metadata.diagrams.plantuml = opts.diagrams
        self.metadata.diagrams.mermaid = opts.diagrams

    def convert(self, markdown_text: str, base_path: str = None) -> bytes:
        front_meta, body = parse_front_matter(markdown_text)
        if front_meta:
            self.metadata = front_meta
        elif self.options.dir:
            self._apply_options()

        tokens = self.parser.parse(body)

        doc_builder = DocumentBuilder(self.metadata)
        self.doc = doc_builder.get_document()

        if self.metadata.diagrams.plantuml or self.metadata.diagrams.mermaid:
            self.diagram_converter = DiagramConverter(self.metadata)

        self._process_tokens(tokens, base_path or os.getcwd())

        self.doc.settings.update_fields = True

        output = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
        self.doc.save(output.name)
        output.close()

        with open(output.name, "rb") as f:
            result = f.read()

        os.unlink(output.name)
        return result

    def _process_tokens(self, tokens: list[Token], base_path: str):
        section_builder = SectionBuilder(self.doc, self.metadata)
        paragraph_builder = ParagraphBuilder(self.doc, self.metadata)
        list_builder = ListBuilder(self.doc, self.metadata)
        code_builder = CodeBuilder(self.doc, self.metadata)
        numbering_builder = NumberingBuilder(self.doc, self.metadata)
        header_footer = HeaderFooterBuilder(self.doc, self.metadata)
        toc_builder = TOCBuilder(self.doc, self.metadata)

        chapter_num = 0
        in_list = False
        list_items = []
        list_ordered = False
        list_level = 0
        current_list_type = None
        skip_until = -1
        i = 0

        if self.metadata.cover_style:
            CoverBuilder(self.doc, self.metadata).build()
            add_page_break(self.doc)
            section_builder.add_new_section()

        if self.metadata.toc.enabled:
            toc_builder.add_toc()

        if self.metadata.lof.enabled:
            add_page_break(self.doc)
            toc_builder.add_lof()

        if self.metadata.lot.enabled:
            add_page_break(self.doc)
            toc_builder.add_lot()

        if self.metadata.toc.enabled or self.metadata.lof.enabled or self.metadata.lot.enabled:
            add_page_break(self.doc)
            section_builder.add_new_section()
            section = self.doc.sections[-1]
            if self.metadata.page_numbers.enabled:
                header_footer.setup_section_numbering(
                    section, self.metadata.page_numbers.body, 1
                )
                header_footer.add_page_numbers(
                    len(self.doc.sections) - 1, self.metadata.page_numbers.body, 1
                )

        has_cover = bool(self.metadata.cover_style)
        has_front = self.metadata.toc.enabled or self.metadata.lof.enabled or self.metadata.lot.enabled
        if has_cover or has_front:
            section = self.doc.sections[0]
            if self.metadata.page_numbers.enabled:
                header_footer.setup_section_numbering(section, "lowerRoman", 1)
                header_footer.add_page_numbers(0, "roman", 1)

        if self.metadata.page_numbers.enabled:
            if not (has_cover or has_front):
                header_footer.add_page_numbers(0, "decimal", 1)

        self.doc.settings.update_fields = True

        while i < len(tokens):
            token = tokens[i]

            if i <= skip_until:
                i += 1
                continue

            if is_page_break_token(token):
                add_page_break(self.doc)
                i += 1
                continue

            if token.type == "heading_open":
                level = int(token.tag[1])
                if level == 1:
                    chapter_num += 1

                i += 1
                inline_text = ""
                while i < len(tokens) and tokens[i].type != "heading_close":
                    if tokens[i].type == "inline":
                        inline_text = tokens[i].content
                    i += 1

                heading_text = inline_text
                paragraph = self.doc.add_paragraph(style=f"Heading {level}")
                dir = detect_text_direction(heading_text)
                if dir == "rtl" or self.metadata.dir == "rtl":
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

                run = paragraph.add_run(heading_text)
                run.font.bold = True
                if dir == "rtl":
                    run.font.name = self.metadata.fonts.arabic
                else:
                    run.font.name = self.metadata.fonts.latin

                if self.metadata.heading_numbering.enabled:
                    numbering_builder.set_heading_number(paragraph, level)

                i += 1
                continue

            if token.type == "paragraph_open":
                i += 1
                inline_text = ""
                while i < len(tokens) and tokens[i].type != "paragraph_close":
                    if tokens[i].type == "inline":
                        inline_text = tokens[i].content
                    i += 1

                if inline_text.strip():
                    paragraph_builder.add_paragraph(inline_text)

                i += 1
                continue

            if token.type == "inline" and token.content.strip():
                paragraph_builder.add_paragraph(token.content)

            if token.type == "fence":
                code = token.content
                language = token.info.strip().split()[0].lower() if token.info else ""

                if language in ("plantuml", "puml") and self.metadata.diagrams.plantuml and self.diagram_converter:
                    img_path = self.diagram_converter.convert_plantuml(code)
                    if img_path:
                        self.figure_counter["figure"] += 1
                        caption = f"{self.metadata.captions.figure_prefix} {chapter_num}-{self.figure_counter['figure']}: PlantUML Diagram"
                        ImageBuilder(self.doc, self.metadata, chapter_num,
                                    figure_counter=self.figure_counter).add_image_direct(img_path, caption)
                        i += 1
                        continue

                if language in ("mermaid", "mmd") and self.metadata.diagrams.mermaid and self.diagram_converter:
                    img_path = self.diagram_converter.convert_mermaid(code)
                    if img_path:
                        self.figure_counter["figure"] += 1
                        caption = f"{self.metadata.captions.figure_prefix} {chapter_num}-{self.figure_counter['figure']}: Mermaid Diagram"
                        ImageBuilder(self.doc, self.metadata, chapter_num,
                                    figure_counter=self.figure_counter).add_image_direct(img_path, caption)
                        i += 1
                        continue

                code_builder.add_code_block(code, language)

            if token.type == "blockquote_open":
                i += 1
                quote_text = ""
                while i < len(tokens) and tokens[i].type != "blockquote_close":
                    if tokens[i].type == "inline":
                        quote_text += tokens[i].content + " "
                    i += 1

                paragraph = self.doc.add_paragraph()
                pf = paragraph.paragraph_format
                pf.left_indent = Cm(1.0)
                pf.space_before = Pt(6)
                pf.space_after = Pt(6)

                pPr = paragraph._element.get_or_add_pPr()
                pBdr = OxmlElement("w:pBdr")
                left = OxmlElement("w:left")
                left.set(qn("w:val"), "single")
                left.set(qn("w:sz"), "12")
                left.set(qn("w:space"), "8")
                left.set(qn("w:color"), "BBBBBB")
                pBdr.append(left)
                pPr.append(pBdr)

                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "F5F5F5")
                shd.set(qn("w:val"), "clear")
                pPr.append(shd)

                run = paragraph.add_run(quote_text.strip())
                run.font.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                dir = detect_text_direction(quote_text)
                if dir == "rtl":
                    run.font.name = self.metadata.fonts.arabic
                else:
                    run.font.name = self.metadata.fonts.latin

                i += 1
                continue

            if token.type == "bullet_list_open":
                list_ordered = False
                in_list = True
                current_list_type = "bullet"
                list_level = 0
                list_items = []
                i += 1
                continue

            if token.type == "ordered_list_open":
                list_ordered = True
                in_list = True
                current_list_type = "ordered"
                list_level = 0
                list_items = []
                i += 1
                continue

            if token.type in ("bullet_list_close", "ordered_list_close"):
                if list_items:
                    for item in list_items:
                        text = item.get("text", "")
                        level = item.get("level", 0)
                        checked = item.get("checked")
                        task = item.get("task", False)
                        num = item.get("num", 1)
                        if list_ordered:
                            dir = detect_text_direction(text)
                            is_arabic = dir == "rtl"
                            paragraph = self.doc.add_paragraph(style="List Paragraph")
                            pf = paragraph.paragraph_format
                            pf.left_indent = Cm(1.27 + level * 0.85)
                            if is_arabic or self.metadata.dir == "rtl":
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                            if task:
                                symbol = "\u2611" if checked else "\u2610"
                                run = paragraph.add_run(f"{symbol} ")
                            else:
                                run = paragraph.add_run()
                                if is_arabic:
                                    run.text = f"{convert_to_arabic_indic(str(num))} - "
                                else:
                                    run.text = f"{num}. "
                            run.font.size = Pt(self.metadata.fonts.size)
                            content_run = paragraph.add_run(text)
                            content_run.font.size = Pt(self.metadata.fonts.size)
                            if is_arabic:
                                content_run.font.name = self.metadata.fonts.arabic
                            else:
                                content_run.font.name = self.metadata.fonts.latin
                        else:
                            paragraph = self.doc.add_paragraph(style="List Paragraph")
                            pf = paragraph.paragraph_format
                            pf.left_indent = Cm(1.27 + level * 0.85)
                            dir = detect_text_direction(text)
                            if dir == "rtl" or self.metadata.dir == "rtl":
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                            bullet = ["●", "○", "▪", "▹"][min(level, 3)]
                            if task:
                                symbol = "\u2611" if checked else "\u2610"
                                run = paragraph.add_run(f"{symbol} ")
                            else:
                                run = paragraph.add_run(f"{bullet} ")
                            run.font.size = Pt(self.metadata.fonts.size)
                            content_run = paragraph.add_run(text)
                            content_run.font.size = Pt(self.metadata.fonts.size)
                            if dir == "rtl":
                                content_run.font.name = self.metadata.fonts.arabic
                            else:
                                content_run.font.name = self.metadata.fonts.latin

                in_list = False
                list_items = []
                i += 1
                continue

            if token.type == "list_item_open":
                item_level = 0
                item_text = ""
                item_checked = None
                item_task = False
                item_num = 1
                i += 1

                while i < len(tokens) and tokens[i].type != "list_item_close":
                    if tokens[i].type == "inline":
                        content = tokens[i].content
                        if content.startswith("[ ] "):
                            item_task = True
                            item_checked = False
                            content = content[4:]
                        elif content.startswith("[x] ") or content.startswith("[X] "):
                            item_task = True
                            item_checked = True
                            content = content[4:]
                        item_text += content
                    i += 1

                list_items.append({
                    "text": item_text.strip(),
                    "level": item_level,
                    "checked": item_checked,
                    "task": item_task,
                    "num": len(list_items) + 1,
                })
                continue

            if token.type == "table_open":
                table_rows = []
                table_headers = []
                in_header = True
                i += 1

                while i < len(tokens) and tokens[i].type != "table_close":
                    if tokens[i].type == "thead_open":
                        in_header = True
                    if tokens[i].type == "tbody_open":
                        in_header = False
                    if tokens[i].type == "tr_open":
                        row_cells = []
                        i += 1
                        while i < len(tokens) and tokens[i].type != "tr_close":
                            if tokens[i].type == "th_open" or tokens[i].type == "td_open":
                                i += 1
                                cell_text = ""
                                while i < len(tokens) and tokens[i].type not in ("th_close", "td_close"):
                                    if tokens[i].type == "inline":
                                        cell_text += tokens[i].content
                                    i += 1
                                row_cells.append(cell_text.strip())
                            i += 1
                        if in_header:
                            table_headers = row_cells
                        else:
                            table_rows.append(row_cells)
                    i += 1

                table_builder = TableBuilder(self.doc, self.metadata, chapter_num,
                                            self.figure_counter)
                table_builder.add_table(table_rows, table_headers if table_headers else None)
                i += 1
                continue

            if token.type == "hr":
                paragraph = self.doc.add_paragraph()
                pf = paragraph.paragraph_format
                pf.space_before = Pt(6)
                pf.space_after = Pt(6)
                pPr = paragraph._element.get_or_add_pPr()
                pBdr = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6")
                bottom.set(qn("w:space"), "1")
                bottom.set(qn("w:color"), "AAAAAA")
                pBdr.append(bottom)
                pPr.append(pBdr)

            if token.type == "image":
                src = token.attrs.get("src", "") if token.attrs else ""
                alt = token.content
                title = token.attrs.get("title", "") if token.attrs else ""
                caption = title or alt
                if src:
                    ImageBuilder(self.doc, self.metadata, chapter_num,
                                base_path=base_path,
                                figure_counter=self.figure_counter).add_image(src, alt, caption)

            i += 1

        if self.diagram_converter:
            self.diagram_converter.cleanup()