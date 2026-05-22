import unicodedata
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from ..models import DocumentMetadata
from ..styles import Fonts, Colors, Spacing


RTL_CHAR_RANGES = [
    (0x0600, 0x06FF),
    (0x0750, 0x077F),
    (0x08A0, 0x08FF),
    (0xFB50, 0xFDFF),
    (0xFE70, 0xFEFF),
    (0x0590, 0x05FF),
    (0x0700, 0x074F),
    (0x0780, 0x07BF),
    (0x07C0, 0x07FF),
    (0x0800, 0x083F),
    (0x0840, 0x085F),
]

ARABIC_INDIC_DIGITS = set("٠١٢٣٤٥٦٧٨٩")
WESTERN_DIGITS = set("0123456789")


def is_rtl_char(char: str) -> bool:
    code = ord(char)
    for start, end in RTL_CHAR_RANGES:
        if start <= code <= end:
            return True
    return False


def detect_text_direction(text: str) -> str:
    rtl_count = sum(1 for c in text if is_rtl_char(c))
    if rtl_count == 0:
        return "ltr"
    total_printable = sum(1 for c in text if not c.isspace())
    if total_printable == 0:
        return "ltr"
    return "rtl" if (rtl_count / total_printable) > 0.3 else "ltr"


def segment_by_direction(text: str) -> list[tuple[str, str]]:
    if not text:
        return [("", "ltr")]
    segments = []
    current_segment = ""
    current_dir = detect_text_direction(text[0]) if text else "ltr"

    for char in text:
        char_dir = detect_text_direction(char) if not char.isspace() else current_dir
        if char.isspace():
            current_segment += char
        elif char_dir == current_dir:
            current_segment += char
        else:
            if current_segment:
                segments.append((current_segment, current_dir))
            current_segment = char
            current_dir = char_dir

    if current_segment:
        segments.append((current_segment, current_dir))

    if not segments:
        segments = [("", "ltr")]

    return segments


def convert_to_arabic_indic(text: str) -> str:
    result = []
    for char in text:
        if char in WESTERN_DIGITS:
            result.append(chr(ord(char) + 0x0660 - ord("0")))
        else:
            result.append(char)
    return "".join(result)


class ParagraphBuilder:
    def __init__(self, doc: Document, metadata: DocumentMetadata):
        self.doc = doc
        self.metadata = metadata

    def add_paragraph(self, text: str, style: str = "Normal", bold: bool = False,
                     italic: bool = False, alignment: str = None) -> "ParagraphBuilder":
        if not text.strip():
            return self

        paragraph = self.doc.add_paragraph(style=style)
        dir = detect_text_direction(text)
        self._apply_direction(paragraph, dir)

        if alignment:
            self._set_alignment(paragraph, alignment)
        elif self.metadata.dir == "rtl":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

        self._add_text_run(paragraph, text, bold=bold, italic=italic)
        return self

    def add_rich_paragraph(self, segments: list[dict]) -> "ParagraphBuilder":
        if not segments:
            return self

        paragraph = self.doc.add_paragraph(style="Normal")
        full_text = "".join(s.get("text", "") for s in segments)
        dir = detect_text_direction(full_text)
        self._apply_direction(paragraph, dir)

        if self.metadata.dir == "rtl":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        for segment in segments:
            text = segment.get("text", "")
            if not text:
                continue
            is_bold = segment.get("bold", False)
            is_italic = segment.get("italic", False)
            is_code = segment.get("code", False)
            is_strikethrough = segment.get("strikethrough", False)
            link = segment.get("link", "")
            color = segment.get("color", None)

            run = paragraph.add_run(text)
            run.font.bold = is_bold
            run.font.italic = is_italic
            run.font.strike = is_strikethrough
            if color:
                run.font.color.rgb = color
            if is_code:
                self._apply_code_format(run)
            else:
                self._apply_segment_font(run, text)
            if link:
                self._make_hyperlink(paragraph, text, link)

        return self

    def _add_text_run(self, paragraph, text: str, bold: bool = False, italic: bool = False):
        directions = segment_by_direction(text)
        for seg_text, seg_dir in directions:
            if not seg_text.strip():
                paragraph.add_run(seg_text)
                continue
            run = paragraph.add_run(seg_text)
            run.font.bold = bold
            run.font.italic = italic
            self._apply_run_direction(run, seg_dir)
            self._apply_segment_font(run, seg_text, seg_dir)

    def _apply_segment_font(self, run, text: str, seg_dir: str = None):
        if seg_dir is None:
            seg_dir = detect_text_direction(text)
        if seg_dir == "rtl":
            run.font.name = self.metadata.fonts.arabic
            run._element.rPr.rFonts.set(qn("w:eastAsia"), self.metadata.fonts.arabic)
        else:
            run.font.name = self.metadata.fonts.latin

    def _apply_code_format(self, run):
        run.font.name = self.metadata.fonts.code
        run.font.size = Pt(9)
        run.font.color.rgb = Colors.CODE_FG

    def _apply_direction(self, paragraph, dir: str):
        pPr = paragraph._element.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        if dir == "rtl":
            bidi.set(qn("w:val"), "1")
        else:
            bidi.set(qn("w:val"), "0")
        pPr.append(bidi)

    def _apply_run_direction(self, run, dir: str):
        rPr = run._element.get_or_add_rPr()
        rtl = OxmlElement("w:rtl")
        if dir == "rtl":
            rtl.set(qn("w:val"), "1")
        else:
            rtl.set(qn("w:val"), "0")
        rPr.append(rtl)

    def _set_alignment(self, paragraph, alignment: str):
        align_map = {
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        paragraph.alignment = align_map.get(alignment, WD_ALIGN_PARAGRAPH.RIGHT)

    def _make_hyperlink(self, paragraph, text: str, url: str):
        part = paragraph.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)

        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        hyperlink.set(qn("w:history"), "1")

        new_run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        c = OxmlElement("w:color")
        c.set(qn("w:val"), "0563C1")
        rPr.append(c)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
        new_run.append(rPr)
        new_run.text = text
        hyperlink.append(new_run)

        run_elem = paragraph.runs[-1]._element
        run_elem.addnext(hyperlink)