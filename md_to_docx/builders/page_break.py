import re
from docx import Document
from docx.enum.text import WD_BREAK


PAGE_BREAK_PATTERNS = [
    re.compile(r"<!--\s*pagebreak\s*-->", re.IGNORECASE),
    re.compile(r"\\pagebreak"),
    re.compile(r"\\newpage"),
    re.compile(r"<div[^>]*style=[\"']page-break-after:\s*always[\"'][^>]*>"),
]


def is_page_break_token(token) -> bool:
    if token.type in ("container_pagebreak_open", "html_block"):
        for pattern in PAGE_BREAK_PATTERNS:
            if pattern.search(token.content):
                return True
    if token.type == "inline":
        content = token.content.strip()
        for pattern in PAGE_BREAK_PATTERNS:
            if pattern.search(content):
                return True
    return False


def clean_page_break_text(text: str) -> str:
    for pattern in PAGE_BREAK_PATTERNS:
        text = pattern.sub("", text)
    return text.strip()


def add_page_break(document: Document):
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)
    return paragraph


def add_page_break_with_paragraph(document: Document, text: str = ""):
    paragraph = document.add_paragraph()
    if text:
        run = paragraph.add_run(text)
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)
    return paragraph