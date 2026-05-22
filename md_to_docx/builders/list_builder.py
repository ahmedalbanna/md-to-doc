from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from ..models import DocumentMetadata
from ..styles import Spacing, Fonts, Colors
from .paragraph_builder import detect_text_direction, convert_to_arabic_indic


ARABIC_INDIC_NUMBERS = ["٠", "١", "٢", "٣", "٤", "٥", "٦", "٧", "٨", "٩"]


class ListBuilder:
    def __init__(self, doc: Document, metadata: DocumentMetadata):
        self.doc = doc
        self.metadata = metadata
        self.list_counters: dict[str, int] = {}
        self.list_levels: dict[int, int] = {}

    def add_bullet_list(self, items: list[dict], level: int = 0) -> "ListBuilder":
        for item in items:
            self._add_list_item(item.get("text", ""), level, ordered=False,
                                checked=item.get("checked"), task=item.get("task", False))
            if "children" in item:
                self.add_bullet_list(item["children"], level + 1)
        return self

    def add_ordered_list(self, items: list[dict], level: int = 0, start: int = 1) -> "ListBuilder":
        counter = start
        for item in items:
            self._add_list_item(item.get("text", ""), level, ordered=True,
                                number=counter, checked=item.get("checked"),
                                task=item.get("task", False))
            if "children" in item:
                self.add_ordered_list(item["children"], level + 1)
            counter += 1
        return self

    def add_list_item(self, text: str, level: int = 0, ordered: bool = False,
                     number: int = 1, checked: bool = None, task: bool = False) -> "ListBuilder":
        self._add_list_item(text, level, ordered, number, checked, task)
        return self

    def _add_list_item(self, text: str, level: int, ordered: bool,
                       number: int = 1, checked: bool = None, task: bool = False):
        paragraph = self.doc.add_paragraph(style="List Paragraph")
        dir = detect_text_direction(text)
        is_arabic = dir == "rtl"

        pf = paragraph.paragraph_format
        indent = Spacing.LIST_INDENT + Cm(level * 0.85)
        pf.left_indent = indent
        pf.space_before = Pt(1)
        pf.space_after = Pt(3)
        pf.tab_stops.add_tab_stop(indent + Cm(0.5))

        if is_arabic or self.metadata.dir == "rtl":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

        if task:
            symbol = "\u2611" if checked else "\u2610"
            run = paragraph.add_run(f"{symbol} ")
        else:
            run = paragraph.add_run()

        if ordered:
            if is_arabic:
                num_text = convert_to_arabic_indic(str(number))
                run.text = f"{num_text} - "
            else:
                run.text = f"{number}. "
        else:
            bullet = ["●", "○", "▪", "▹", "◆", "▸", "•"][min(level, 6)]
            run.text = f"{bullet} "

        run.font.size = Pt(self.metadata.fonts.size)
        if is_arabic:
            run.font.name = self.metadata.fonts.arabic
        else:
            run.font.name = self.metadata.fonts.latin

        content_run = paragraph.add_run(text)
        content_run.font.size = Pt(self.metadata.fonts.size)
        if is_arabic:
            content_run.font.name = self.metadata.fonts.arabic
        else:
            content_run.font.name = self.metadata.fonts.latin

        self._apply_rtl_to_paragraph(paragraph, is_arabic)

    def _apply_rtl_to_paragraph(self, paragraph, is_rtl: bool):
        pPr = paragraph._element.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1" if is_rtl else "0")
        pPr.append(bidi)