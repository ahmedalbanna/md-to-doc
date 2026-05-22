from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from ..models import DocumentMetadata
from ..styles import Colors, Fonts, Spacing


class CodeBuilder:
    def __init__(self, doc: Document, metadata: DocumentMetadata):
        self.doc = doc
        self.metadata = metadata

    def add_code_block(self, code: str, language: str = "") -> "CodeBuilder":
        paragraph = self.doc.add_paragraph()
        pf = paragraph.paragraph_format
        pf.space_before = Pt(8)
        pf.space_after = Pt(8)
        pf.left_indent = Spacing.CODE_LEFT_INDENT
        pf.right_indent = Spacing.CODE_RIGHT_INDENT

        self._set_background(paragraph)

        run = paragraph.add_run(code)
        run.font.name = self.metadata.fonts.code
        run.font.size = Pt(9)
        run.font.color.rgb = Colors.CODE_FG

        if language:
            self._add_language_label(language)

        return self

    def add_inline_code(self, text: str) -> "CodeBuilder":
        paragraph = self.doc.add_paragraph()
        run = paragraph.add_run(text)
        run.font.name = self.metadata.fonts.code
        run.font.size = Pt(9)
        run.font.color.rgb = Colors.CODE_FG
        return self

    def _set_background(self, paragraph):
        pPr = paragraph._element.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F2F2F2")
        shd.set(qn("w:val"), "clear")
        pPr.append(shd)

    def _add_language_label(self, language: str):
        paragraph = self.doc.add_paragraph()
        pf = paragraph.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)
        pf.left_indent = Spacing.CODE_LEFT_INDENT

        run = paragraph.add_run(language)
        run.font.size = Pt(7)
        run.font.italic = True
        run.font.color.rgb = Colors.SECONDARY
        run.font.name = self.metadata.fonts.latin