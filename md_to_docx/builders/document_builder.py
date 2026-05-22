from docx import Document
from docx.shared import Cm, Pt
from docx.enum.section import WD_ORIENT
from ..models import DocumentMetadata
from ..styles import PageSetup, Fonts, Colors


class DocumentBuilder:
    def __init__(self, metadata: DocumentMetadata):
        self.metadata = metadata
        self.doc = Document()

    def setup_document_defaults(self):
        style = self.doc.styles["Normal"]
        font = style.font
        font.name = self.metadata.fonts.latin
        font.size = Pt(self.metadata.fonts.size)
        font.color.rgb = Colors.BODY

        rtl_run = style.font
        rtl_run.name = self.metadata.fonts.arabic

        pf = style.paragraph_format
        pf.space_before = Pt(3)
        pf.space_after = Pt(6)

    def setup_sections(self):
        section = self.doc.sections[0]
        section.page_width = PageSetup.PAGE_WIDTH
        section.page_height = PageSetup.PAGE_HEIGHT
        section.top_margin = PageSetup.MARGIN_TOP
        section.bottom_margin = PageSetup.MARGIN_BOTTOM
        section.left_margin = PageSetup.MARGIN_LEFT
        section.right_margin = PageSetup.MARGIN_RIGHT
        section.header_distance = PageSetup.HEADER_DISTANCE
        section.footer_distance = PageSetup.FOOTER_DISTANCE

        self._setup_rtl(section)

    def _setup_rtl(self, section):
        sect_pr = section._sectPr
        if self.metadata.dir == "rtl":
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            bidi = OxmlElement("w:bidi")
            bidi.set(qn("w:val"), "1")
            sect_pr.append(bidi)

            props = sect_pr.find(qn("w:docGrid"))
            if props is None:
                from lxml import etree
                props = OxmlElement("w:docGrid")
                props.set(qn("w:linePitch"), "360")
                sect_pr.append(props)

    def add_heading_style(self, level: int):
        style_name = f"Heading {level}"
        if style_name in self.doc.styles:
            style = self.doc.styles[style_name]
        else:
            style = self.doc.styles.add_style(style_name, 1)

        style.font.name = self.metadata.fonts.heading
        style.font.color.rgb = Colors.HEADING
        style.font.bold = True
        size_map = {1: 26, 2: 22, 3: 18, 4: 16, 5: 14, 6: 12}
        style.font.size = Pt(size_map.get(level, 12))

        pf = style.paragraph_format
        pf.space_before = Pt(18 if level <= 2 else 12)
        pf.space_after = Pt(8)
        pf.keep_with_next = True
        pf.keep_together = True

    def setup_styles(self):
        self.setup_document_defaults()
        for level in range(1, 7):
            self.add_heading_style(level)

    def get_document(self) -> Document:
        self.setup_document_defaults()
        self.setup_sections()
        self.setup_styles()
        return self.doc