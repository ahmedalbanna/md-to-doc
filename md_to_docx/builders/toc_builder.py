from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from ..models import DocumentMetadata
from ..styles import Colors
from .paragraph_builder import detect_text_direction


class TOCBuilder:
    def __init__(self, doc: Document, metadata: DocumentMetadata):
        self.doc = doc
        self.metadata = metadata

    def add_toc(self, title: str = None, max_level: int = 3) -> "TOCBuilder":
        toc_title = title or self.metadata.toc.title
        self._add_toc_heading(toc_title)
        self._insert_toc_field(max_level)
        self._add_empty_line()
        return self

    def add_lof(self, title: str = None) -> "TOCBuilder":
        lof_title = title or self.metadata.lof.title
        self._add_toc_heading(lof_title)
        self._insert_lof_field()
        self._add_empty_line()
        return self

    def add_lot(self, title: str = None) -> "TOCBuilder":
        lot_title = title or self.metadata.lot.title
        self._add_toc_heading(lot_title)
        self._insert_lot_field()
        self._add_empty_line()
        return self

    def _add_toc_heading(self, title: str):
        paragraph = self.doc.add_paragraph()
        if self.metadata.dir == "rtl":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

        run = paragraph.add_run(title)
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = Colors.HEADING
        dir = detect_text_direction(title)
        if dir == "rtl":
            run.font.name = self.metadata.fonts.arabic
        else:
            run.font.name = self.metadata.fonts.latin

        pf = paragraph.paragraph_format
        pf.space_before = Pt(24)
        pf.space_after = Pt(12)
        pf.keep_with_next = True

    def _insert_toc_field(self, max_level: int = 3):
        paragraph = self.doc.add_paragraph()
        pf = paragraph.paragraph_format
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)

        run = paragraph.add_run()

        instr_text = f' TOC \\o "1-{max_level}" \\h \\z \\u '

        fld_char_begin = OxmlElement("w:fldChar")
        fld_char_begin.set(qn("w:fldCharType"), "begin")
        run._element.append(fld_char_begin)

        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = instr_text
        run._element.append(instr)

        fld_char_separate = OxmlElement("w:fldChar")
        fld_char_separate.set(qn("w:fldCharType"), "separate")
        run._element.append(fld_char_separate)

        run2 = paragraph.add_run("[سيتم تحديث جدول المحتويات عند الفتح في Word]")
        run2.font.size = Pt(10)
        run2.font.color.rgb = Colors.SECONDARY
        run2.font.italic = True

        fld_char_end = OxmlElement("w:fldChar")
        fld_char_end.set(qn("w:fldCharType"), "end")
        run._element.append(fld_char_end)

    def _insert_lof_field(self):
        paragraph = self.doc.add_paragraph()
        run = paragraph.add_run()

        fld_char_begin = OxmlElement("w:fldChar")
        fld_char_begin.set(qn("w:fldCharType"), "begin")
        run._element.append(fld_char_begin)

        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = ' TOC \\f \\h \\z '
        run._element.append(instr)

        fld_char_separate = OxmlElement("w:fldChar")
        fld_char_separate.set(qn("w:fldCharType"), "separate")
        run._element.append(fld_char_separate)

        run2 = paragraph.add_run("[سيتم تحديث قائمة الأشكال عند الفتح في Word]")
        run2.font.size = Pt(10)
        run2.font.color.rgb = Colors.SECONDARY
        run2.font.italic = True

        fld_char_end = OxmlElement("w:fldChar")
        fld_char_end.set(qn("w:fldCharType"), "end")
        run._element.append(fld_char_end)

    def _insert_lot_field(self):
        paragraph = self.doc.add_paragraph()
        run = paragraph.add_run()

        fld_char_begin = OxmlElement("w:fldChar")
        fld_char_begin.set(qn("w:fldCharType"), "begin")
        run._element.append(fld_char_begin)

        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = ' TOC \\c "table" \\h \\z '
        run._element.append(instr)

        fld_char_separate = OxmlElement("w:fldChar")
        fld_char_separate.set(qn("w:fldCharType"), "separate")
        run._element.append(fld_char_separate)

        run2 = paragraph.add_run("[سيتم تحديث قائمة الجداول عند الفتح في Word]")
        run2.font.size = Pt(10)
        run2.font.color.rgb = Colors.SECONDARY
        run2.font.italic = True

        fld_char_end = OxmlElement("w:fldChar")
        fld_char_end.set(qn("w:fldCharType"), "end")
        run._element.append(fld_char_end)

    def _add_empty_line(self):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(6)