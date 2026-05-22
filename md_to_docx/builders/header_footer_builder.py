from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from ..models import DocumentMetadata
from ..styles import Colors


class HeaderFooterBuilder:
    def __init__(self, doc: Document, metadata: DocumentMetadata):
        self.doc = doc
        self.metadata = metadata

    def add_page_numbers(self, section_index: int = 0, format_type: str = "decimal",
                         start_number: int = 1) -> "HeaderFooterBuilder":
        section = self.doc.sections[section_index]

        footer = section.footer
        footer.is_linked_to_previous = False

        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()

        if self.metadata.dir == "rtl":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        pf = paragraph.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)

        run = paragraph.add_run()
        run.font.size = Pt(9)
        run.font.color.rgb = Colors.SECONDARY

        fldChar1 = OxmlElement("w:fldChar")
        fldChar1.set(qn("w:fldCharType"), "begin")
        run._element.append(fldChar1)

        instrText = OxmlElement("w:instrText")
        instrText.set(qn("xml:space"), "preserve")

        fmt_map = {
            "western": "",
            "arabic-indic": r"\* Arabic-Indic",
            "roman": r"\* Roman",
            "roman-upper": r"\* UpperRoman",
        }
        fmt = fmt_map.get(format_type, "")
        instrText.text = f" PAGE {fmt}"
        run._element.append(instrText)

        fldChar2 = OxmlElement("w:fldChar")
        fldChar2.set(qn("w:fldCharType"), "end")
        run._element.append(fldChar2)

        pgNumType = section._sectPr.find(qn("w:pgNumType"))
        if pgNumType is None:
            pgNumType = OxmlElement("w:pgNumType")
            section._sectPr.append(pgNumType)
        pgNumType.set(qn("w:start"), str(start_number))

        return self

    def add_header_text(self, text: str, section_index: int = 0) -> "HeaderFooterBuilder":
        section = self.doc.sections[section_index]
        header = section.header
        header.is_linked_to_previous = False

        paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(text)
        run.font.size = Pt(8)
        run.font.color.rgb = Colors.SECONDARY
        run.font.italic = True

        if self.metadata.dir == "rtl":
            run.font.name = self.metadata.fonts.arabic
        else:
            run.font.name = self.metadata.fonts.latin

        return self

    def setup_section_numbering(self, section, format_type: str, start_number: int = 1):
        sect_pr = section._sectPr
        pgNumType = sect_pr.find(qn("w:pgNumType"))
        if pgNumType is None:
            pgNumType = OxmlElement("w:pgNumType")
            sect_pr.append(pgNumType)

        fmt_map = {
            "western": "decimal",
            "arabic-indic": "decimal",
            "roman": "lowerRoman",
            "roman-upper": "upperRoman",
        }
        fmt = fmt_map.get(format_type, "decimal")
        pgNumType.set(qn("w:fmt"), fmt)
        pgNumType.set(qn("w:start"), str(start_number))

    def setup_bidi_footer(self, section):
        footer = section.footer
        for paragraph in footer.paragraphs:
            pPr = paragraph._element.get_or_add_pPr()
            bidi = OxmlElement("w:bidi")
            bidi.set(qn("w:val"), "1")
            pPr.append(bidi)

    def remove_header_footer_link(self, section):
        header = section.header
        footer = section.footer
        header.is_linked_to_previous = False
        footer.is_linked_to_previous = False