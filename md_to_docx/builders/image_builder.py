import os
import tempfile
import urllib.request
import urllib.parse
from pathlib import Path
from docx import Document
from docx.shared import Inches, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from ..models import DocumentMetadata
from ..styles import Spacing, Fonts, Colors
from .paragraph_builder import detect_text_direction


class ImageBuilder:
    def __init__(self, doc: Document, metadata: DocumentMetadata, chapter_num: int = 0,
                 base_path: str = None, figure_counter: dict = None):
        self.doc = doc
        self.metadata = metadata
        self.chapter_num = chapter_num
        self.base_path = base_path or os.getcwd()
        self.figure_counter = figure_counter or {"figure": 0, "table": 0}

    def add_image(self, src: str, alt: str = "", caption: str = None,
                  width: float = 5.5) -> "ImageBuilder":
        image_path = self._resolve_image_path(src)
        if not image_path:
            alt_paragraph = self.doc.add_paragraph()
            alt_run = alt_paragraph.add_run(f"[Image: {alt or src}]")
            alt_run.font.italic = True
            alt_run.font.color.rgb = Colors.SECONDARY
            return self

        paragraph = self.doc.add_paragraph()
        if self.metadata.dir == "rtl":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        try:
            run = paragraph.add_run()
            run.add_picture(image_path, width=Inches(min(width, 6.0)))
        except Exception:
            return self

        if caption:
            self.figure_counter["figure"] += 1
            caption_text = self._build_caption(caption)
            self._add_caption(caption_text)

        return self

    def add_image_direct(self, image_path: str, caption: str = None,
                         width: float = 5.5) -> "ImageBuilder":
        if not os.path.exists(image_path):
            return self

        paragraph = self.doc.add_paragraph()
        if self.metadata.dir == "rtl":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        try:
            run = paragraph.add_run()
            run.add_picture(image_path, width=Inches(min(width, 6.0)))
        except Exception:
            return self

        if caption:
            self.figure_counter["figure"] += 1
            caption_text = self._build_caption(caption)
            self._add_caption(caption_text)

        return self

    def _resolve_image_path(self, src: str) -> str:
        if not src:
            return ""
        if urllib.parse.urlparse(src).scheme in ("http", "https"):
            try:
                ext = os.path.splitext(src)[1] or ".png"
                tmp = tempfile.NamedTemporaryFile(delete=False, suffix=ext)
                urllib.request.urlretrieve(src, tmp.name)
                return tmp.name
            except Exception:
                return ""
        path = Path(src)
        if path.is_absolute():
            return str(path) if path.exists() else ""
        full_path = Path(self.base_path) / path
        return str(full_path) if full_path.exists() else ""

    def _build_caption(self, caption: str) -> str:
        prefix = self.metadata.captions.figure_prefix
        num = self.figure_counter.get("figure", 1)
        if self.metadata.captions.by_chapter:
            return f"{prefix} {self.chapter_num}-{num}: {caption}"
        return f"{prefix} {num}: {caption}"

    def _add_caption(self, text: str):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Spacing.CAPTION_BEFORE
        paragraph.paragraph_format.space_after = Spacing.CAPTION_AFTER

        if self.metadata.dir == "rtl":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        run = paragraph.add_run(text)
        run.font.size = Spacing.SIZE_CAPTION
        run.font.bold = True
        dir = detect_text_direction(text)
        if dir == "rtl":
            run.font.name = self.metadata.fonts.arabic
        else:
            run.font.name = self.metadata.fonts.latin
        run.font.color.rgb = Colors.PRIMARY