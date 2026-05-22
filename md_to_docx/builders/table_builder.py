from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from ..models import DocumentMetadata
from ..styles import Colors, Fonts, Spacing
from .paragraph_builder import detect_text_direction


class TableBuilder:
    def __init__(self, doc: Document, metadata: DocumentMetadata, chapter_num: int = 0,
                 table_counter: dict = None):
        self.doc = doc
        self.metadata = metadata
        self.chapter_num = chapter_num
        self.table_counter = table_counter or {"table": 0, "figure": 0}

    def add_table(self, rows: list[list[str]], headers: list[str] = None,
                  caption: str = None, alignment: list[str] = None) -> "TableBuilder":
        num_rows = len(rows) + (1 if headers else 0)
        num_cols = max((len(headers) if headers else 0),
                       max((len(r) for r in rows), default=0))

        if num_cols == 0:
            return self

        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        self._set_table_borders(table)
        self._apply_table_direction(table)

        if headers:
            self._add_header_row(table, headers, num_cols)

        self._add_data_rows(table, rows, headers is not None, num_cols)

        self._auto_fit_columns(table, num_cols)

        if caption:
            self.table_counter["table"] += 1
            caption_text = self._build_caption(caption, "table")
            self._add_caption(caption_text)

        return self

    def _add_header_row(self, table, headers: list[str], num_cols: int):
        header_row = table.rows[0]
        for i in range(min(len(headers), num_cols)):
            cell = header_row.cells[i]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            self._format_header_cell(cell, paragraph, headers[i])

    def _format_header_cell(self, cell, paragraph, text: str):
        cell._element.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "1B4F72")
        shading.set(qn("w:val"), "clear")
        cell._element.tcPr.append(shading)

        run = paragraph.add_run(text)
        run.font.color.rgb = Colors.TABLE_HEADER_FG
        run.font.bold = True
        run.font.size = Pt(11)
        dir = detect_text_direction(text)
        if dir == "rtl":
            run.font.name = self.metadata.fonts.arabic
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            run.font.name = self.metadata.fonts.latin
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(4)

    def _add_data_rows(self, table, rows: list[list[str]], has_header: bool, num_cols: int):
        start_row = 1 if has_header else 0
        for row_idx, row_data in enumerate(rows):
            row = table.rows[row_idx + start_row]
            for col_idx in range(num_cols):
                cell = row.cells[col_idx]
                cell.text = ""
                paragraph = cell.paragraphs[0]
                text = row_data[col_idx] if col_idx < len(row_data) else ""

                if row_idx % 2 == 1:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "F8F8F8")
                    shading.set(qn("w:val"), "clear")
                    cell._element.tcPr.append(shading)

                run = paragraph.add_run(text)
                run.font.size = Pt(10)
                dir = detect_text_direction(text)
                if dir == "rtl":
                    run.font.name = self.metadata.fonts.arabic
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    run.font.name = self.metadata.fonts.latin

                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)

    def _set_table_borders(self, table):
        tbl = table._tbl
        tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
        borders = OxmlElement("w:tblBorders")
        for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            border = OxmlElement(f"w:{border_name}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:space"), "0")
            border.set(qn("w:color"), "CCCCCC")
            borders.append(border)
        tblPr.append(borders)

    def _apply_table_direction(self, table):
        if self.metadata.dir == "rtl":
            tbl = table._tbl
            tblPr = tbl.tblPr
            if tblPr is None:
                tblPr = OxmlElement("w:tblPr")
                tbl.insert(0, tblPr)
            bidi = OxmlElement("w:bidiVisual")
            bidi.set(qn("w:val"), "1")
            tblPr.append(bidi)

    def _auto_fit_columns(self, table, num_cols: int):
        tbl = table._tbl
        tblPr = tbl.tblPr
        if tblPr is None:
            tblPr = OxmlElement("w:tblPr")
            tbl.insert(0, tblPr)
        tblW = OxmlElement("w:tblW")
        tblW.set(qn("w:w"), "5000")
        tblW.set(qn("w:type"), "pct")
        tblPr.append(tblW)

    def _build_caption(self, caption: str, cap_type: str) -> str:
        prefix = self.metadata.captions.table_prefix if cap_type == "table" else self.metadata.captions.figure_prefix
        num = self.table_counter.get(cap_type, 1)
        if self.metadata.captions.by_chapter:
            return f"{prefix} {self.chapter_num}-{num}: {caption}"
        return f"{prefix} {num}: {caption}"

    def _add_caption(self, text: str):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Spacing.CAPTION_BEFORE
        paragraph.paragraph_format.space_after = Spacing.CAPTION_AFTER

        if self.metadata.dir == "rtl":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        run = paragraph.add_run(text)
        run.font.size = Spacing.SIZE_CAPTION
        run.font.bold = True
        dir = detect_text_direction(text)
        if dir == "rtl":
            run.font.name = self.metadata.fonts.arabic
        else:
            run.font.name = self.metadata.fonts.latin
        run.font.color.rgb = Colors.PRIMARY