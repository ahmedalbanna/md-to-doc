import os
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from ..models import DocumentMetadata
from ..styles import Colors, Fonts, Spacing, CoverStyles
from .paragraph_builder import detect_text_direction


class CoverBuilder:
    def __init__(self, doc: Document, metadata: DocumentMetadata):
        self.doc = doc
        self.metadata = metadata

    def build(self) -> "CoverBuilder":
        style = self.metadata.cover_style or CoverStyles.CLASSIC
        method_map = {
            CoverStyles.CLASSIC: self._build_classic,
            CoverStyles.MODERN: self._build_modern,
            CoverStyles.MINIMAL: self._build_minimal,
            CoverStyles.ELEGANT: self._build_elegant,
        }
        builder = method_map.get(style, self._build_classic)
        builder()
        return self

    def _build_classic(self):
        self._add_vertical_space(3)
        self._add_centered_title(self.metadata.title or "عنوان المستند")
        self._add_horizontal_line(Colors.COVER_ACCENT)
        if self.metadata.subtitle:
            self._add_centered_subtitle(self.metadata.subtitle)
        self._add_vertical_space(4)
        self._add_centered_meta(self.metadata.author, self.metadata.date, self.metadata.version)
        self._add_logo()

    def _build_modern(self):
        self._add_side_bar()
        self._add_vertical_space(4)
        self._add_title_with_accent(self.metadata.title or "عنوان المستند")
        if self.metadata.subtitle:
            self._add_centered_subtitle(self.metadata.subtitle, Colors.SECONDARY)
        self._add_vertical_space(1)
        self._add_horizontal_line(Colors.PRIMARY)
        self._add_vertical_space(3)
        self._add_centered_meta(self.metadata.author, self.metadata.date, self.metadata.version)
        self._add_logo()

    def _build_minimal(self):
        self._add_vertical_space(5)
        self._add_thin_line()
        self._add_vertical_space(1)
        self._add_centered_title(self.metadata.title or "عنوان المستند", Fonts.SIZE_COVER_TITLE)
        if self.metadata.subtitle:
            self._add_centered_subtitle(self.metadata.subtitle, Colors.SECONDARY)
        self._add_vertical_space(1)
        self._add_thin_line()
        self._add_vertical_space(4)
        self._add_centered_meta(self.metadata.author, self.metadata.date)

    def _build_elegant(self):
        self._add_ornamental_border()
        self._add_vertical_space(3)
        self._add_ornamental_line()
        self._add_vertical_space(1)
        self._add_centered_title(self.metadata.title or "عنوان المستند", Fonts.SIZE_COVER_TITLE,
                                 Colors.COVER_GOLD)
        if self.metadata.subtitle:
            self._add_centered_subtitle(self.metadata.subtitle, Colors.COVER_GOLD)
        self._add_vertical_space(1)
        self._add_ornamental_line()
        self._add_vertical_space(3)
        self._add_centered_meta(self.metadata.author, self.metadata.date, self.metadata.version)
        self._add_logo()

    def _add_centered_title(self, text: str, size=Fonts.SIZE_COVER_TITLE, color=Colors.HEADING):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.size = size
        run.font.bold = True
        run.font.color.rgb = color
        dir = detect_text_direction(text)
        run.font.name = self.metadata.fonts.arabic if dir == "rtl" else self.metadata.fonts.latin
        paragraph.paragraph_format.space_after = Pt(6)

    def _add_centered_subtitle(self, text: str, color=Colors.SECONDARY):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.size = Fonts.SIZE_COVER_SUBTITLE
        run.font.italic = True
        run.font.color.rgb = color
        dir = detect_text_direction(text)
        run.font.name = self.metadata.fonts.arabic if dir == "rtl" else self.metadata.fonts.latin
        paragraph.paragraph_format.space_after = Pt(12)

    def _add_centered_meta(self, author: str = "", date: str = "", version: str = ""):
        parts = [p for p in [author, date, version] if p]
        if not parts:
            return
        text = "  |  ".join(parts)
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.size = Fonts.SIZE_COVER_META
        run.font.color.rgb = Colors.SECONDARY
        dir = detect_text_direction(text)
        run.font.name = self.metadata.fonts.arabic if dir == "rtl" else self.metadata.fonts.latin

    def _add_horizontal_line(self, color: RGBColor):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = paragraph.paragraph_format
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)

        pPr = paragraph._element.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "1")
        hex_color = "{:02X}{:02X}{:02X}".format(color[0], color[1], color[2])
        bottom.set(qn("w:color"), hex_color)
        pBdr.append(bottom)
        pPr.append(pBdr)

    def _add_thin_line(self):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = paragraph._element.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "999999")
        pBdr.append(bottom)
        pPr.append(pBdr)

    def _add_ornamental_line(self):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("─" * 40 + " ✦ " + "─" * 40)
        run.font.color.rgb = Colors.COVER_GOLD
        run.font.size = Pt(14)

    def _add_ornamental_border(self):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = paragraph._element.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        for side in ["top", "bottom"]:
            border = OxmlElement(f"w:{side}")
            border.set(qn("w:val"), "double")
            border.set(qn("w:sz"), "8")
            border.set(qn("w:space"), "4")
            border.set(qn("w:color"), "D4A537")
            pBdr.append(border)
        pPr.append(pBdr)
        paragraph.paragraph_format.space_before = Pt(20)
        paragraph.paragraph_format.space_after = Pt(20)

    def _add_side_bar(self):
        for _ in range(8):
            paragraph = self.doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = paragraph.add_run(" " * 5)
            run.font.size = Pt(14)
            self._add_run_shading(run, "1B4F72")

    def _add_run_shading(self, run, color: str):
        rPr = run._element.get_or_add_rPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), color)
        shd.set(qn("w:val"), "clear")
        rPr.append(shd)

    def _add_title_with_accent(self, text: str):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run(text)
        run.font.size = Fonts.SIZE_COVER_TITLE
        run.font.bold = True
        run.font.color.rgb = Colors.PRIMARY
        dir = detect_text_direction(text)
        run.font.name = self.metadata.fonts.arabic if dir == "rtl" else self.metadata.fonts.latin

    def _add_vertical_space(self, lines: int = 1):
        for _ in range(lines):
            paragraph = self.doc.add_paragraph()
            run = paragraph.add_run("")
            run.font.size = Pt(12)

    def _add_logo(self):
        if not self.metadata.logo:
            return
        logo_path = self.metadata.logo
        if not os.path.exists(logo_path):
            return
        try:
            paragraph = self.doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(logo_path, width=Inches(2.0))
        except Exception:
            pass